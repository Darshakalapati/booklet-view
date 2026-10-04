"""
Reading an uploaded Excel (.xlsx) or CSV file into tables.

The first row that has anything in it becomes the column names. Every row
after it becomes a row of the table. Each column gets a type by looking at
its values:

    all numbers                      -> number
    all real dates / YYYY-MM-DD      -> date
    all yes/no or true/false         -> checkbox
    anything else                    -> text

An Excel file with several sheets gives one table per sheet.

Nothing here touches the database; ``read_upload`` only returns plain data,
and the view saves it.
"""

import csv
import datetime
import io
from dataclasses import dataclass, field

from django.conf import settings

from .values import clean_value, is_empty

YES_NO_WORDS = {"true", "false", "yes", "no"}


class ImportProblem(Exception):
    """Something is wrong with the uploaded file; the message is shown to the user."""


@dataclass
class ParsedTable:
    name: str
    columns: list = field(default_factory=list)   # [(name, data_type), ...]
    rows: list = field(default_factory=list)      # [[cell, cell, ...], ...] cleaned, in column order


def read_upload(uploaded_file):
    """Return a list of ``ParsedTable`` for an uploaded .xlsx or .csv file."""
    filename = (uploaded_file.name or "").strip()
    lowered = filename.lower()
    if uploaded_file.size > settings.IMPORT_MAX_FILE_BYTES:
        limit_mb = settings.IMPORT_MAX_FILE_BYTES // (1024 * 1024)
        raise ImportProblem(f"The file is larger than {limit_mb} MB.")

    base_name = filename.rsplit(".", 1)[0].strip() or "Imported table"
    if lowered.endswith(".csv"):
        grids = [(base_name, read_csv_grid(uploaded_file.read()))]
    elif lowered.endswith(".xlsx"):
        grids = read_xlsx_grids(uploaded_file, base_name)
    elif lowered.endswith(".xls"):
        raise ImportProblem("Old .xls files are not supported. Save the file as .xlsx or .csv and upload it again.")
    else:
        raise ImportProblem("Only Excel (.xlsx) and CSV (.csv) files can be uploaded.")

    tables = [parse_grid(name, grid) for name, grid in grids]
    tables = [table for table in tables if table is not None]
    if not tables:
        raise ImportProblem("The file has no data in it.")
    return tables


def read_csv_grid(raw_bytes):
    try:
        text = raw_bytes.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = raw_bytes.decode("latin-1")
    try:
        dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    return [list(row) for row in csv.reader(io.StringIO(text, newline=""), dialect)]


def read_xlsx_grids(uploaded_file, base_name):
    from openpyxl import load_workbook
    from openpyxl.utils.exceptions import InvalidFileException
    from zipfile import BadZipFile

    try:
        workbook = load_workbook(uploaded_file, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, ValueError):
        raise ImportProblem("The Excel file could not be opened. It may be damaged or not a real .xlsx file.")
    try:
        sheets = [sheet for sheet in workbook.worksheets if sheet.sheet_state == "visible"]
        grids = []
        for sheet in sheets:
            name = sheet.title if len(sheets) > 1 else base_name
            grids.append((name, [list(row) for row in sheet.iter_rows(values_only=True)]))
        return grids
    finally:
        workbook.close()


def parse_grid(name, grid):
    """Turn a grid of raw cells into a ``ParsedTable``; ``None`` if it is empty."""
    grid = [row for row in grid if any(not is_empty(cell) for cell in row)]
    if not grid:
        return None

    width = max(len(row) for row in grid)
    grid = [list(row) + [None] * (width - len(row)) for row in grid]

    # Drop columns that are empty all the way down (common in spreadsheets).
    used = [index for index in range(width) if any(not is_empty(row[index]) for row in grid)]
    grid = [[row[index] for index in used] for row in grid]

    header, body = grid[0], grid[1:]
    if len(header) > settings.IMPORT_MAX_COLUMNS:
        raise ImportProblem(f'"{name}" has more than {settings.IMPORT_MAX_COLUMNS} columns.')
    if len(body) > settings.IMPORT_MAX_ROWS:
        raise ImportProblem(f'"{name}" has more than {settings.IMPORT_MAX_ROWS} rows.')

    names = unique_names(header)
    table = ParsedTable(name=name[:200])
    cleaned_columns = []
    for index, column_name in enumerate(names):
        raw_values = [row[index] for row in body]
        data_type = guess_type(raw_values)
        table.columns.append((column_name, data_type))
        cleaned_columns.append([clean_value(data_type, value) for value in raw_values])
    table.rows = [list(cells) for cells in zip(*cleaned_columns)] if cleaned_columns else []
    return table


def unique_names(header):
    """Column names from the header row: no blanks and no duplicates."""
    names, seen = [], {}
    for index, cell in enumerate(header, start=1):
        name = "" if is_empty(cell) else str(cell).strip()
        name = (name or f"Column {index}")[:190]
        seen[name.lower()] = seen.get(name.lower(), 0) + 1
        if seen[name.lower()] > 1:
            name = f"{name} ({seen[name.lower()]})"
        names.append(name)
    return names


def guess_type(raw_values):
    values = [value for value in raw_values if not is_empty(value)]
    if not values:
        return "text"
    for data_type, accepts in (("checkbox", looks_like_checkbox), ("number", looks_like_number), ("date", looks_like_date)):
        if all(accepts(value) for value in values):
            return data_type
    return "text"


def looks_like_checkbox(value):
    if isinstance(value, bool):
        return True
    return isinstance(value, str) and value.strip().lower() in YES_NO_WORDS


def looks_like_number(value):
    if isinstance(value, bool):
        return False
    if isinstance(value, (int, float)):
        return True
    if isinstance(value, str):
        text = value.strip()
        # Keep things like phone numbers and codes with leading zeros as text.
        if len(text) > 1 and text[0] == "0" and text[1].isdigit():
            return False
        if text.startswith("+"):
            return False
        try:
            clean_value("number", text)
            return True
        except ValueError:
            return False
    return False


def looks_like_date(value):
    if isinstance(value, (datetime.date, datetime.datetime)):
        return True
    if isinstance(value, str):
        try:
            datetime.date.fromisoformat(value.strip())
            return True
        except ValueError:
            return False
    return False
